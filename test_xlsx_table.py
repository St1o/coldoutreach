import io
import unittest

import openpyxl

import xlsx_table

HEADER = ["Название", "ФИО директора", "Почта", "Основной ОКВЭД"]


def read(rows, header=HEADER):
    return openpyxl.load_workbook(io.BytesIO(xlsx_table.build(header, rows))).active


class XlsxTableTest(unittest.TestCase):
    def test_header_rows_and_layout(self):
        sheet = read([['ООО "РОМАШКА"', "Иванов Иван", "a@b.ru", "Разработка ПО"], ["ООО Б", "", "c@d.ru", ""]])
        self.assertEqual(sheet.title, "Компании")
        self.assertEqual([[cell.value for cell in row] for row in sheet.iter_rows()],
                         [HEADER, ['ООО "РОМАШКА"', "Иванов Иван", "a@b.ru", "Разработка ПО"],
                          ["ООО Б", None, "c@d.ru", None]])             # пустые значения Excel читает как пустые ячейки
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertGreater(sheet.column_dimensions["A"].width, len('ООО "РОМАШКА"'))

    def test_no_rows_still_gives_a_file_with_the_header(self):
        sheet = read([])
        self.assertEqual([cell.value for cell in sheet[1]], HEADER)
        self.assertEqual(sheet.max_row, 1)

    def test_a_formula_looking_value_stays_text(self):
        sheet = read([['=HYPERLINK("http://x.example","Тут")', "+7 123", "@x", "-1"]])
        for cell in sheet[2]:
            self.assertEqual(cell.data_type, "s", cell.value)
        self.assertEqual(sheet["A2"].value, '=HYPERLINK("http://x.example","Тут")')

    def test_control_characters_are_removed(self):
        self.assertEqual(read([["ООО\x00 \x0bА", "", "", ""]])["A2"].value, "ООО А")


if __name__ == "__main__":
    unittest.main()
