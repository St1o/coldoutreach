"""Excel-файл (.xlsx) с таблицей выгрузки: открывается в Excel и загружается в Google Таблицы
(Файл -> Импортировать -> Загрузить).

Названия и ФИО приходят из внешних источников, поэтому все значения записываются как текст:
строка вида «=ФОРМУЛА(...)» не станет формулой. Служебные символы, которые не допускает формат, убираются.
"""
import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SHEET_TITLE = "Компании"
MAX_COLUMN_WIDTH = 60


def build(header: list, rows: list) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    widths = [len(title) for title in header]

    def put(row_number, values, bold=False):
        for column, value in enumerate(values, 1):
            text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
            cell = sheet.cell(row=row_number, column=column, value=text)
            cell.data_type = "s"                    # только текст, даже если значение начинается с «=»
            if bold:
                cell.font = Font(bold=True)
            widths[column - 1] = max(widths[column - 1], len(text))

    put(1, header, bold=True)
    for number, row in enumerate(rows, 2):
        put(number, row)
    sheet.freeze_panes = "A2"
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = min(width + 2, MAX_COLUMN_WIDTH)

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
